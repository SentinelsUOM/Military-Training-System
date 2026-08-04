# -*- coding: utf-8 -*-
"""
Renders briefing_content.py's BLOCKS into three synchronized outputs:
    MODULE2_EVALUATION_BRIEFING.md
    MODULE2_EVALUATION_BRIEFING.docx
    MODULE2_EVALUATION_BRIEFING.pdf

This is a LIVING DOCUMENT generator. To update the briefing: edit BLOCKS in
briefing_content.py, then re-run this script. All three outputs stay in sync
because they're all rendered from the same source list — there is no separate
prose to maintain per format.
"""
import os
from briefing_content import BLOCKS, LAST_UPDATED

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
OUT_MD   = os.path.join(ROOT, 'MODULE2_EVALUATION_BRIEFING.md')
OUT_DOCX = os.path.join(ROOT, 'MODULE2_EVALUATION_BRIEFING.docx')
OUT_PDF  = os.path.join(ROOT, 'MODULE2_EVALUATION_BRIEFING.pdf')

INK_HEX   = '0b0b0b'
INK2_HEX  = '52514e'
MUTED_HEX = '898781'
ACCENT_HEX = '2a78d6'
GRID_HEX  = 'e1e0d9'


# ════════════════════════════════════════════════════════════════════════════
# Markdown renderer
# ════════════════════════════════════════════════════════════════════════════

def render_md(blocks, path):
    lines = []
    for b in blocks:
        kind = b[0]
        if kind == 'h1':
            lines.append(f'# {b[1]}\n')
        elif kind == 'h2':
            lines.append(f'## {b[1]}\n')
        elif kind == 'h3':
            lines.append(f'### {b[1]}\n')
        elif kind == 'p':
            lines.append(f'{b[1]}\n')
        elif kind == 'bp':
            lines.append(f'**{b[1]}**\n')
        elif kind == 'quote':
            quoted = '\n'.join('> ' + ln for ln in b[1].split('\n'))
            lines.append(f'{quoted}\n')
        elif kind == 'bullet':
            lines.extend(f'- {item}' for item in b[1])
            lines.append('')
        elif kind == 'numbered':
            lines.extend(f'{i+1}. {item}' for i, item in enumerate(b[1]))
            lines.append('')
        elif kind == 'table':
            headers, rows = b[1], b[2]
            lines.append('| ' + ' | '.join(headers) + ' |')
            lines.append('|' + '|'.join(['---'] * len(headers)) + '|')
            for row in rows:
                lines.append('| ' + ' | '.join(str(c) for c in row) + ' |')
            lines.append('')
        elif kind == 'image':
            fname, caption = b[1], b[2]
            lines.append(f'![{caption}](Report_Figures/{fname})')
            lines.append(f'*{caption}*\n')
        elif kind == 'hr':
            lines.append('---\n')
        elif kind == 'pagebreak':
            pass  # no-op in markdown
    with open(path, 'w', encoding='utf-8') as f:
        f.write('\n'.join(lines))
    print('Wrote', path)


# ════════════════════════════════════════════════════════════════════════════
# DOCX renderer
# ════════════════════════════════════════════════════════════════════════════

def render_docx(blocks, path):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    INK   = RGBColor(0x0b, 0x0b, 0x0b)
    INK2  = RGBColor(0x52, 0x51, 0x4e)
    MUTED = RGBColor(0x89, 0x87, 0x81)
    HEADER_FILL = 'DCE6F5'

    def set_cell_shading(cell, hex_color):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), hex_color)
        tcPr.append(shd)

    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Segoe UI'
    style.font.size = Pt(10.5)
    style.font.color.rgb = INK
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Inches(0.9)

    for b in blocks:
        kind = b[0]
        if kind == 'h1':
            h = doc.add_heading(b[1], level=1)
            h.runs[0].font.color.rgb = INK
        elif kind == 'h2':
            h = doc.add_heading(b[1], level=2)
            h.runs[0].font.color.rgb = INK
        elif kind == 'h3':
            h = doc.add_heading(b[1], level=3)
            h.runs[0].font.color.rgb = INK
        elif kind == 'p':
            doc.add_paragraph(b[1])
        elif kind == 'bp':
            p = doc.add_paragraph()
            r = p.add_run(b[1])
            r.bold = True
        elif kind == 'quote':
            for ln in b[1].split('\n'):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.35)
                r = p.add_run(ln)
                r.italic = True
                r.font.color.rgb = INK2
                r.font.name = 'Consolas' if any(c in ln for c in ('=', '{', '}')) else 'Segoe UI'
        elif kind == 'bullet':
            for item in b[1]:
                doc.add_paragraph(item, style='List Bullet')
        elif kind == 'numbered':
            for item in b[1]:
                doc.add_paragraph(item, style='List Number')
        elif kind == 'table':
            headers, rows = b[1], b[2]
            table = doc.add_table(rows=1, cols=len(headers))
            table.style = 'Table Grid'
            hdr = table.rows[0].cells
            for i, htext in enumerate(headers):
                hdr[i].text = htext
                set_cell_shading(hdr[i], HEADER_FILL)
                for p in hdr[i].paragraphs:
                    for r in p.runs:
                        r.bold = True
                        r.font.size = Pt(9.5)
            for row in rows:
                cells = table.add_row().cells
                for i, val in enumerate(row):
                    cells[i].text = str(val)
                    for p in cells[i].paragraphs:
                        for r in p.runs:
                            r.font.size = Pt(9.5)
            doc.add_paragraph()
        elif kind == 'image':
            fname, caption = b[1], b[2]
            img_path = os.path.join(HERE, fname)
            if os.path.exists(img_path):
                doc.add_picture(img_path, width=Inches(6.0))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap = doc.add_paragraph(caption)
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in cap.runs:
                r.italic = True
                r.font.size = Pt(9)
                r.font.color.rgb = MUTED
            doc.add_paragraph()
        elif kind == 'hr':
            p = doc.add_paragraph('─' * 60)
            p.runs[0].font.color.rgb = MUTED
        elif kind == 'pagebreak':
            doc.add_page_break()

    doc.save(path)
    print('Wrote', path)


# ════════════════════════════════════════════════════════════════════════════
# PDF renderer (reportlab platypus)
# ════════════════════════════════════════════════════════════════════════════

def render_pdf(blocks, path):
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.units import inch
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, ListFlowable,
                                     ListItem, Table, TableStyle, Image, PageBreak, HRFlowable)
    from reportlab.lib.enums import TA_LEFT

    INK   = colors.HexColor('#0b0b0b')
    INK2  = colors.HexColor('#52514e')
    MUTED = colors.HexColor('#898781')
    GRID  = colors.HexColor('#e1e0d9')
    HEADER_FILL = colors.HexColor('#DCE6F5')

    styles = getSampleStyleSheet()
    h1 = ParagraphStyle('H1', parent=styles['Heading1'], fontName='Helvetica-Bold',
                         fontSize=16, textColor=INK, spaceBefore=16, spaceAfter=8)
    h2 = ParagraphStyle('H2', parent=styles['Heading2'], fontName='Helvetica-Bold',
                         fontSize=13, textColor=INK, spaceBefore=12, spaceAfter=6)
    h3 = ParagraphStyle('H3', parent=styles['Heading3'], fontName='Helvetica-Bold',
                         fontSize=11, textColor=INK, spaceBefore=10, spaceAfter=5)
    body = ParagraphStyle('Body', parent=styles['Normal'], fontName='Helvetica',
                           fontSize=10, leading=14, textColor=INK, spaceAfter=8,
                           alignment=TA_LEFT)
    bold_body = ParagraphStyle('BoldBody', parent=body, fontName='Helvetica-Bold')
    quote_style = ParagraphStyle('Quote', parent=body, fontName='Helvetica-Oblique',
                                  textColor=INK2, leftIndent=18, fontSize=9.5, leading=13)
    caption_style = ParagraphStyle('Caption', parent=body, fontName='Helvetica-Oblique',
                                    fontSize=8.5, textColor=MUTED, alignment=1,
                                    spaceAfter=12)
    cell_style = ParagraphStyle('Cell', parent=body, fontSize=8.5, leading=11, spaceAfter=0)
    cell_head_style = ParagraphStyle('CellHead', parent=cell_style, fontName='Helvetica-Bold')

    def esc(text):
        return (text.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;'))

    story = []
    for b in blocks:
        kind = b[0]
        if kind == 'h1':
            story.append(Paragraph(esc(b[1]), h1))
        elif kind == 'h2':
            story.append(Paragraph(esc(b[1]), h2))
        elif kind == 'h3':
            story.append(Paragraph(esc(b[1]), h3))
        elif kind == 'p':
            story.append(Paragraph(esc(b[1]), body))
        elif kind == 'bp':
            story.append(Paragraph(esc(b[1]), bold_body))
        elif kind == 'quote':
            for ln in b[1].split('\n'):
                story.append(Paragraph(esc(ln) if ln.strip() else '&nbsp;', quote_style))
            story.append(Spacer(1, 8))
        elif kind == 'bullet':
            items = [ListItem(Paragraph(esc(i), body), leftIndent=14) for i in b[1]]
            story.append(ListFlowable(items, bulletType='bullet', start='•',
                                       leftIndent=18, spaceBefore=2, spaceAfter=8))
        elif kind == 'numbered':
            items = [ListItem(Paragraph(esc(i), body), leftIndent=14) for i in b[1]]
            story.append(ListFlowable(items, bulletType='1', leftIndent=18,
                                       spaceBefore=2, spaceAfter=8))
        elif kind == 'table':
            headers, rows = b[1], b[2]
            data = [[Paragraph(esc(str(h)), cell_head_style) for h in headers]]
            for row in rows:
                data.append([Paragraph(esc(str(c)), cell_style) for c in row])
            n_cols = len(headers)
            avail_width = LETTER[0] - 1.6 * inch
            col_width = avail_width / n_cols
            t = Table(data, colWidths=[col_width] * n_cols, repeatRows=1)
            t.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), HEADER_FILL),
                ('GRID', (0, 0), (-1, -1), 0.5, GRID),
                ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                ('TOPPADDING', (0, 0), (-1, -1), 4),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
                ('LEFTPADDING', (0, 0), (-1, -1), 5),
                ('RIGHTPADDING', (0, 0), (-1, -1), 5),
            ]))
            story.append(t)
            story.append(Spacer(1, 10))
        elif kind == 'image':
            fname, caption = b[1], b[2]
            img_path = os.path.join(HERE, fname)
            if os.path.exists(img_path):
                img = Image(img_path, width=5.6 * inch, height=5.6 * inch * 0.55)
                img.hAlign = 'CENTER'
                story.append(img)
            story.append(Paragraph(esc(caption), caption_style))
        elif kind == 'hr':
            story.append(Spacer(1, 4))
            story.append(HRFlowable(width='100%', color=MUTED, thickness=0.5))
            story.append(Spacer(1, 10))
        elif kind == 'pagebreak':
            story.append(PageBreak())

    doc = SimpleDocTemplate(path, pagesize=LETTER,
                             leftMargin=0.8 * inch, rightMargin=0.8 * inch,
                             topMargin=0.8 * inch, bottomMargin=0.8 * inch,
                             title='Module 2 — Evaluation Briefing')
    doc.build(story)
    print('Wrote', path)


if __name__ == '__main__':
    render_md(BLOCKS, OUT_MD)
    render_docx(BLOCKS, OUT_DOCX)
    render_pdf(BLOCKS, OUT_PDF)
    print(f'\nAll three outputs regenerated (content last updated {LAST_UPDATED}).')
