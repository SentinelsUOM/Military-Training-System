# -*- coding: utf-8 -*-
"""
Parses MODULE1_EVALUATOR_QA_SCRIPT.md directly (the single source of truth —
no re-transcription, so there's zero risk of drifting from the actual content)
and renders it to .docx and .pdf, matching the visual style used for the other
Module docs in this repo (generate_briefing.py's palette).

Handles: h1/h2/h3 headers, paragraphs, bullet lists, numbered lists, tables,
code fences (```...```), blockquotes (> ...), horizontal rules, and inline
**bold**/*italic* emphasis (rendered as real bold/italic runs, not literal
asterisks).
"""
import os
import re

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
SRC = os.path.join(ROOT, 'MODULE1_EVALUATOR_QA_SCRIPT.md')
OUT_DOCX = os.path.join(ROOT, 'MODULE1_EVALUATOR_QA_SCRIPT.docx')
OUT_PDF = os.path.join(ROOT, 'MODULE1_EVALUATOR_QA_SCRIPT.pdf')

INK_HEX, INK2_HEX, MUTED_HEX = '0b0b0b', '52514e', '898781'
HEADER_FILL_HEX = 'DCE6F5'

# ── Inline emphasis tokenizer (**bold** / *italic*) ─────────────────────────
_TOKEN_RE = re.compile(r'(\*\*[^*]+?\*\*|\*[^*]+?\*)')


def tokenize_inline(text):
    """Split into (content, bold, italic) runs."""
    parts = _TOKEN_RE.split(text)
    tokens = []
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            tokens.append((part[2:-2], True, False))
        elif part.startswith('*') and part.endswith('*'):
            tokens.append((part[1:-1], False, True))
        else:
            tokens.append((part, False, False))
    return tokens


# ── Markdown -> block parser ─────────────────────────────────────────────────

def parse_markdown(text):
    lines = text.split('\n')
    blocks = []
    i, n = 0, len(lines)
    para_buf = []

    def flush_para():
        if para_buf:
            blocks.append(('p', ' '.join(para_buf).strip()))
            para_buf.clear()

    while i < n:
        raw = lines[i]
        stripped = raw.strip()

        # code fence
        if stripped.startswith('```'):
            flush_para()
            code_lines = []
            i += 1
            while i < n and not lines[i].strip().startswith('```'):
                code_lines.append(lines[i])
                i += 1
            i += 1  # skip closing fence
            blocks.append(('code', '\n'.join(code_lines)))
            continue

        # horizontal rule
        if stripped == '---':
            flush_para()
            blocks.append(('hr', None))
            i += 1
            continue

        # headers
        m = re.match(r'^(#{1,3})\s+(.*)', stripped)
        if m:
            flush_para()
            kind = {1: 'h1', 2: 'h2', 3: 'h3'}[len(m.group(1))]
            blocks.append((kind, m.group(2).strip()))
            i += 1
            continue

        # table
        if stripped.startswith('|'):
            flush_para()
            table_lines = []
            while i < n and lines[i].strip().startswith('|'):
                table_lines.append(lines[i].strip())
                i += 1
            headers = [c.strip() for c in table_lines[0].strip('|').split('|')]
            rows = [[c.strip() for c in rl.strip('|').split('|')] for rl in table_lines[2:]]
            blocks.append(('table', (headers, rows)))
            continue

        # bullet list
        if stripped.startswith('- '):
            flush_para()
            items = []
            while i < n and lines[i].strip().startswith('- '):
                items.append(lines[i].strip()[2:].strip())
                i += 1
            blocks.append(('bullet', items))
            continue

        # numbered list
        if re.match(r'^\d+\.\s+', stripped):
            flush_para()
            items = []
            while i < n and re.match(r'^\d+\.\s+', lines[i].strip()):
                items.append(re.sub(r'^\d+\.\s+', '', lines[i].strip()))
                i += 1
            blocks.append(('numbered', items))
            continue

        # blockquote
        if stripped.startswith('>'):
            flush_para()
            q_lines = []
            while i < n and lines[i].strip().startswith('>'):
                q = lines[i].strip()[1:].strip()
                q_lines.append(q)
                i += 1
            blocks.append(('quote', ' '.join(q_lines)))
            continue

        # whole-line bold (e.g. **SAY THIS:**)
        if stripped.startswith('**') and stripped.endswith('**') and stripped.count('**') == 2:
            flush_para()
            blocks.append(('bp', stripped.strip('*')))
            i += 1
            continue

        # blank line = paragraph break
        if stripped == '':
            flush_para()
            i += 1
            continue

        para_buf.append(stripped)
        i += 1

    flush_para()
    return blocks


# ── DOCX renderer ─────────────────────────────────────────────────────────

def render_docx(blocks, path):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    INK = RGBColor(0x0b, 0x0b, 0x0b)
    INK2 = RGBColor(0x52, 0x51, 0x4e)
    MUTED = RGBColor(0x89, 0x87, 0x81)

    def shade(cell, hexcolor):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), hexcolor)
        tcPr.append(shd)

    def add_rich(paragraph, text, color=None, size=None):
        for content, bold, italic in tokenize_inline(text):
            r = paragraph.add_run(content)
            r.bold = bold
            r.italic = italic
            if color:
                r.font.color.rgb = color
            if size:
                r.font.size = size

    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Segoe UI'
    style.font.size = Pt(10.5)
    style.font.color.rgb = INK
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Inches(0.9)

    for kind, val in blocks:
        if kind == 'h1':
            h = doc.add_heading('', level=1)
            add_rich(h, val, color=INK)
        elif kind == 'h2':
            h = doc.add_heading('', level=2)
            add_rich(h, val, color=INK)
        elif kind == 'h3':
            h = doc.add_heading('', level=3)
            add_rich(h, val, color=INK)
        elif kind == 'p':
            p = doc.add_paragraph()
            add_rich(p, val)
        elif kind == 'bp':
            p = doc.add_paragraph()
            r = p.add_run(val)
            r.bold = True
        elif kind == 'quote':
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.35)
            add_rich(p, val, color=INK2)
            for r in p.runs:
                r.italic = True
        elif kind == 'code':
            for line in val.split('\n'):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.35)
                r = p.add_run(line if line.strip() else ' ')
                r.font.name = 'Consolas'
                r.font.size = Pt(9)
                r.font.color.rgb = INK2
        elif kind == 'bullet':
            for item in val:
                p = doc.add_paragraph(style='List Bullet')
                add_rich(p, item)
        elif kind == 'numbered':
            for item in val:
                p = doc.add_paragraph(style='List Number')
                add_rich(p, item)
        elif kind == 'table':
            headers, rows = val
            table = doc.add_table(rows=1, cols=len(headers))
            table.style = 'Table Grid'
            hdr = table.rows[0].cells
            for i, htext in enumerate(headers):
                hdr[i].text = ''
                p = hdr[i].paragraphs[0]
                add_rich(p, htext)
                for r in p.runs:
                    r.bold = True
                    r.font.size = Pt(9.5)
                shade(hdr[i], HEADER_FILL_HEX)
            for row in rows:
                cells = table.add_row().cells
                for i, valtext in enumerate(row):
                    cells[i].text = ''
                    p = cells[i].paragraphs[0]
                    add_rich(p, valtext)
                    for r in p.runs:
                        r.font.size = Pt(9.5)
            doc.add_paragraph()
        elif kind == 'hr':
            p = doc.add_paragraph('─' * 60)
            p.runs[0].font.color.rgb = MUTED

    doc.save(path)
    print('Wrote', path)


# ── PDF renderer ─────────────────────────────────────────────────────────

def esc_xml(text):
    return text.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')


def to_pdf_markup(text):
    esc_text = esc_xml(text)
    out = []
    for content, bold, italic in tokenize_inline(esc_text):
        if bold:
            out.append(f'<b>{content}</b>')
        elif italic:
            out.append(f'<i>{content}</i>')
        else:
            out.append(content)
    return ''.join(out)


def render_pdf(blocks, path):
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.units import inch
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, ListFlowable,
                                     ListItem, Table, TableStyle, Preformatted, HRFlowable)
    from reportlab.lib.enums import TA_LEFT

    INK = colors.HexColor('#0b0b0b')
    INK2 = colors.HexColor('#52514e')
    MUTED = colors.HexColor('#898781')
    GRID = colors.HexColor('#e1e0d9')
    HEADER_FILL = colors.HexColor('#DCE6F5')

    styles = getSampleStyleSheet()
    h1 = ParagraphStyle('H1', parent=styles['Heading1'], fontName='Helvetica-Bold',
                         fontSize=16, textColor=INK, spaceBefore=16, spaceAfter=8)
    h2 = ParagraphStyle('H2', parent=styles['Heading2'], fontName='Helvetica-Bold',
                         fontSize=13, textColor=INK, spaceBefore=14, spaceAfter=6)
    h3 = ParagraphStyle('H3', parent=styles['Heading3'], fontName='Helvetica-Bold',
                         fontSize=11, textColor=INK, spaceBefore=10, spaceAfter=5)
    body = ParagraphStyle('Body', parent=styles['Normal'], fontName='Helvetica',
                           fontSize=10, leading=14, textColor=INK, spaceAfter=8,
                           alignment=TA_LEFT)
    bold_body = ParagraphStyle('BoldBody', parent=body, fontName='Helvetica-Bold')
    quote_style = ParagraphStyle('Quote', parent=body, fontName='Helvetica-Oblique',
                                  textColor=INK2, leftIndent=18, fontSize=9.5, leading=13)
    code_style = ParagraphStyle('Code', parent=body, fontName='Courier', fontSize=8.3,
                                 leading=11, textColor=INK2, leftIndent=14,
                                 backColor=colors.HexColor('#f5f4ef'), spaceAfter=8,
                                 borderPadding=6)
    cell_style = ParagraphStyle('Cell', parent=body, fontSize=8.5, leading=11, spaceAfter=0)
    cell_head_style = ParagraphStyle('CellHead', parent=cell_style, fontName='Helvetica-Bold')

    story = []
    for kind, val in blocks:
        if kind == 'h1':
            story.append(Paragraph(to_pdf_markup(val), h1))
        elif kind == 'h2':
            story.append(Paragraph(to_pdf_markup(val), h2))
        elif kind == 'h3':
            story.append(Paragraph(to_pdf_markup(val), h3))
        elif kind == 'p':
            story.append(Paragraph(to_pdf_markup(val), body))
        elif kind == 'bp':
            story.append(Paragraph(to_pdf_markup(val), bold_body))
        elif kind == 'quote':
            story.append(Paragraph(to_pdf_markup(val), quote_style))
            story.append(Spacer(1, 4))
        elif kind == 'code':
            story.append(Preformatted(val, code_style))
        elif kind == 'bullet':
            items = [ListItem(Paragraph(to_pdf_markup(it), body), leftIndent=14) for it in val]
            story.append(ListFlowable(items, bulletType='bullet', start='•',
                                       leftIndent=18, spaceBefore=2, spaceAfter=8))
        elif kind == 'numbered':
            items = [ListItem(Paragraph(to_pdf_markup(it), body), leftIndent=14) for it in val]
            story.append(ListFlowable(items, bulletType='1', leftIndent=18,
                                       spaceBefore=2, spaceAfter=8))
        elif kind == 'table':
            headers, rows = val
            data = [[Paragraph(to_pdf_markup(h), cell_head_style) for h in headers]]
            for row in rows:
                data.append([Paragraph(to_pdf_markup(c), cell_style) for c in row])
            n_cols = len(headers)
            avail_width = LETTER[0] - 1.6 * inch
            col_width = avail_width / n_cols
            t = Table(data, colWidths=[col_width] * n_cols, repeatRows=1)
            t.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), HEADER_FILL),
                ('GRID', (0, 0), (-1, -1), 0.5, GRID),
                ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                ('TOPPADDING', (0, 0), (-1, -1), 4),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
                ('LEFTPADDING', (0, 0), (-1, -1), 5),
                ('RIGHTPADDING', (0, 0), (-1, -1), 5),
            ]))
            story.append(t)
            story.append(Spacer(1, 10))
        elif kind == 'hr':
            story.append(Spacer(1, 4))
            story.append(HRFlowable(width='100%', color=MUTED, thickness=0.5))
            story.append(Spacer(1, 10))

    doc = SimpleDocTemplate(path, pagesize=LETTER,
                             leftMargin=0.8 * inch, rightMargin=0.8 * inch,
                             topMargin=0.8 * inch, bottomMargin=0.8 * inch,
                             title='Module 1 — Evaluator Q&A Script')
    doc.build(story)
    print('Wrote', path)


if __name__ == '__main__':
    with open(SRC, 'r', encoding='utf-8') as f:
        text = f.read()
    blocks = parse_markdown(text)
    render_docx(blocks, OUT_DOCX)
    render_pdf(blocks, OUT_PDF)
